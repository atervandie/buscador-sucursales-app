import openpyxl
from supabase import create_client
import os
from dotenv import load_dotenv

load_dotenv()

# Leer Excel
wb = openpyxl.load_workbook('CROISSANT_SUCURSALES_UNIFICADAS_83.xlsx')
ws = wb.active

excel_sucursales = set()
for row in range(2, ws.max_row + 1):
    numero = ws.cell(row=row, column=1).value
    if numero:
        excel_sucursales.add(str(numero).strip())

print(f"Sucursales en Excel: {len(excel_sucursales)}")

# Leer Supabase
supabase = create_client(os.getenv('SUPABASE_URL'), os.getenv('SUPABASE_API_KEY'))
data = supabase.table('sucursales').select('numero_sucursal').execute()
db_sucursales = {item['numero_sucursal'] for item in data.data}

print(f"Sucursales en Supabase: {len(db_sucursales)}")

# Encontrar diferencias
faltantes = excel_sucursales - db_sucursales
extras = db_sucursales - excel_sucursales

if faltantes:
    print(f"\n🔴 NO cargadas en Supabase: {sorted(faltantes)}")
if extras:
    print(f"\n🟢 Cargadas pero NO en Excel: {sorted(extras)}")
if not faltantes and not extras:
    print("\n✅ Todas coinciden perfectamente")