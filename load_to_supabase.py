#!/usr/bin/env python3
"""
Script para cargar datos de Excel a Supabase
Uso: python3 load_to_supabase.py <archivo_excel>
"""

import sys
import os
from dotenv import load_dotenv
from typing import List, Dict  # ← AGREGAR ESTA LÍNEA

# Cargar variables de .env
load_dotenv()

import openpyxl
from supabase import create_client, Client

def load_excel_data(excel_path: str) -> List[Dict]:
    """
    Carga datos del Excel y devuelve lista de diccionarios
    """
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    # Leer headers (fila 1)
    headers = []
    for col in range(1, ws.max_column + 1):
        header = ws.cell(row=1, column=col).value
        if header:
            # Convertir nombres de columnas a snake_case
            header_snake = header.lower().replace(' ', '_').replace('á', 'a').replace('ó', 'o').replace('í', 'i').replace('ú', 'u')
            headers.append(header_snake)
    
    # Leer datos (desde fila 2)
    datos = []
    for row in range(2, ws.max_row + 1):
        fila_data = {}
        for col_idx, header in enumerate(headers, 1):
            value = ws.cell(row=row, column=col_idx).value
            fila_data[header] = value
        
        # Solo agregar si tiene número de sucursal
        if fila_data.get('no._sucursal'):
            # Mapear campos al esquema de Supabase
            registro = {
                'numero_sucursal': str(fila_data.get('no._sucursal', '')).strip(),
                'razon_social': fila_data.get('razon_social', ''),
                'nombre_sucursal': fila_data.get('sucursal', ''),
                'domicilio_fisico': fila_data.get('domicilio_fisico', ''),
                'localidad': fila_data.get('localidad', ''),
                'marca': fila_data.get('marca', ''),
                'geolocalizacion_url': fila_data.get('geolocalizacion', ''),
            }
            datos.append(registro)
    
    return datos

def load_to_supabase(datos: List[Dict], supabase_url: str, supabase_key: str):
    """
    Carga datos a Supabase
    """
    try:
        # Crear cliente de Supabase
        supabase: Client = create_client(supabase_url, supabase_key)
        
        print(f"📌 Conectado a Supabase")
        print(f"📊 Cargando {len(datos)} sucursales...\n")
        
        # Batch insert (de 100 en 100)
        batch_size = 100
        for i in range(0, len(datos), batch_size):
            batch = datos[i:i+batch_size]
            
            # Insertar batch
            resultado = supabase.table('sucursales').upsert(
                batch,
                on_conflict='numero_sucursal'
            ).execute()
            
            print(f"✅ Cargadas {min(i + batch_size, len(datos))}/{len(datos)} sucursales")
        
        print(f"\n🎉 Carga completada exitosamente!")
        print(f"📍 Total de sucursales: {len(datos)}")
        
    except Exception as e:
        print(f"❌ Error al conectar con Supabase: {e}")
        sys.exit(1)

def main():
    # Verificar argumentos
    if len(sys.argv) < 2:
        print("Uso: python3 load_to_supabase.py <archivo_excel>")
        print("\nVariables de entorno requeridas:")
        print("  SUPABASE_URL: URL de tu proyecto Supabase")
        print("  SUPABASE_API_KEY: Clave API de Supabase (Service Role Key)")
        sys.exit(1)
    
    excel_path = sys.argv[1]
    supabase_url = os.getenv('SUPABASE_URL')
    supabase_key = os.getenv('SUPABASE_API_KEY')
    
    if not supabase_url or not supabase_key:
        print("❌ Error: Variables de entorno SUPABASE_URL y SUPABASE_API_KEY no configuradas")
        sys.exit(1)
    
    if not os.path.exists(excel_path):
        print(f"❌ Error: Archivo {excel_path} no encontrado")
        sys.exit(1)
    
    # Cargar datos
    print(f"📖 Leyendo Excel: {excel_path}")
    datos = load_excel_data(excel_path)
    
    # Mostrar preview
    print("\nPreview de datos:")
    for i, d in enumerate(datos[:3]):
        print(f"\n{i+1}. {d['numero_sucursal']} - {d['razon_social'][:50]}")
        print(f"   📍 {d['domicilio_fisico'][:60]}")
        if d['geolocalizacion_url']:
            print(f"   🗺️  Maps: {d['geolocalizacion_url'][:40]}...")
    
    print(f"\n... ({len(datos)} sucursales en total)")
    
    # Cargar a Supabase
    load_to_supabase(datos, supabase_url, supabase_key)

if __name__ == '__main__':
    main()
